import sys
from datetime import date
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import extract  # noqa: E402
import make_sample_invoices  # noqa: E402


@pytest.fixture(scope="module")
def invoices(tmp_path_factory):
    folder = tmp_path_factory.mktemp("invoices")
    make_sample_invoices.main(folder)
    return {inv["file"]: inv for inv in map(extract.read_invoice, sorted(Path(folder).glob("*.pdf")))}


@pytest.mark.parametrize("text, expected", [
    ("Total 34,023.30", 34023.30),
    ("Amount due EGP 23,193.30", 23193.30),
    ("Grand total EGP 8,094.00", 8094.00),
    ("Subtotal", None),
])
def test_parse_money(text, expected):
    assert extract.parse_money(text) == expected


@pytest.mark.parametrize("text, expected", [
    ("04/03/2026", date(2026, 3, 4)),  # day first
    ("2026-03-04", date(2026, 3, 4)),
    ("07 Mar 2026 Sinai Retail Group", date(2026, 3, 7)),  # text after the date is ignored
    ("March 15, 2026", date(2026, 3, 15)),
    ("Mar 15, 2026", date(2026, 3, 15)),
    ("4 March 2026", date(2026, 3, 4)),
    ("no date here", None),
])
def test_parse_date(text, expected):
    assert extract.parse_date(text) == expected


def test_every_layout_is_read(invoices):
    assert {inv["supplier"] for inv in invoices.values() if inv["supplier"]} == {
        "Nile Paper Co.", "Delta Print House", "Canal Tech Services"}
    nile = invoices["nile_NPC-2026-0107.pdf"]
    assert (nile["invoice_no"], nile["date"]) == ("NPC-2026-0107", date(2026, 3, 4))
    assert invoices["delta_DPH-1043.pdf"]["invoice_no"] == "DPH/1043"
    assert invoices["canal_7781.pdf"]["invoice_no"] == "7781"


def test_clean_invoices_add_up(invoices):
    clean = [inv for inv in invoices.values() if not inv["problems"]]
    assert len(clean) == 11
    for inv in clean:
        assert sum(i["amount"] for i in inv["items"]) == pytest.approx(inv["subtotal"])
        assert inv["subtotal"] + inv["vat"] == pytest.approx(inv["total"])
        assert inv["vat"] == pytest.approx(inv["subtotal"] * 0.14, abs=0.01)


def test_wrong_total_is_flagged(invoices):
    problems = invoices["delta_DPH-1045.pdf"]["problems"]
    assert len(problems) == 1 and "+100.00" in problems[0]


def test_scan_is_flagged_not_guessed(invoices):
    scan = invoices["scanned_invoice.pdf"]
    assert scan["problems"] == ["No text layer (scanned image): needs OCR or manual entry"]
    assert scan["total"] is None and scan["items"] == []


def test_check_catches_bad_lines():
    invoice = {"invoice_no": "1", "date": date(2026, 1, 1), "subtotal": 300.0, "vat": 42.0, "total": 342.0,
               "items": [{"description": "Pens", "qty": 2, "unit_price": 50.0, "amount": 120.0},
                         {"description": "Paper", "qty": 1, "unit_price": 150.0, "amount": 150.0}]}
    problems = extract.check(invoice)
    assert any("Pens" in p for p in problems)
    assert any("Lines add up to 270.00" in p for p in problems)


def test_workbook(invoices, tmp_path):
    tables = extract.to_tables(list(invoices.values()))
    extract.write_workbook(tables, tmp_path / "invoices.xlsx")
    wb = openpyxl.load_workbook(tmp_path / "invoices.xlsx")
    assert wb.sheetnames == ["Invoices", "Line items", "Needs review"]
    assert wb["Invoices"].max_row == 14 and wb["Needs review"].max_row == 3
    assert wb["Invoices"]["F2"].number_format == "#,##0.00"
